from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

BASE = Path(__file__).resolve().parent
QA = BASE / '_architecture_qa'
QA.mkdir(exist_ok=True)
OUT = BASE / 'CLSU_System_Architecture.docx'

# A static architecture figure with separate human and agent entry points.
im = Image.new('RGB', (1800, 1500), 'white')
draw = ImageDraw.Draw(im)
fontroot = Path('C:/Windows/Fonts')
def font(n, bold=False):
    return ImageFont.truetype(str(fontroot / ('arialbd.ttf' if bold else 'arial.ttf')), n)
def box(rect, title, lines=(), fill='#F2F6F8'):
    draw.rounded_rectangle(rect, radius=14, fill=fill, outline='#526575', width=3)
    x0,y0,x1,y1 = rect
    draw.text(((x0+x1)/2,y0+20),title,font=font(29,True),fill='#152F42',anchor='mt')
    for i,line in enumerate(lines):
        draw.text(((x0+x1)/2,y0+64+i*37),line,font=font(25),fill='#22313C',anchor='mt')
def arrow(points, label=None):
    draw.line(points,fill='#526575',width=4)
    x,y=points[-1]; px,py=points[-2]
    import math
    a=math.atan2(y-py,x-px)
    tri=[(x,y),(x-17*math.cos(a-.45),y-17*math.sin(a-.45)),(x-17*math.cos(a+.45),y-17*math.sin(a+.45))]
    draw.polygon(tri,fill='#526575')
    if label:
        lx,ly=label[0],label[1]
        draw.text((lx,ly),label[2],font=font(22),fill='#22313C',anchor='mm')
box((40,30,1110,195),'React web application', ['Public feedback | Requester portal | Personnel | MISO', 'Embedded assistant sends chat to the server'])
box((1190,30,1760,195),'External AI agent',['Compatible MCP host', 'Authenticated administrator'])
box((40,290,650,440),'Application API',['Forms, tracking and reports'])
box((740,290,1240,440),'Embedded agent runtime',['Model calls and MCP client'])
box((1330,290,1760,440),'AI provider',['Selected model API'])
box((740,530,1760,660),'Laravel MCP server',['Approved read tools and operational action tools'])
box((40,755,1760,900),'Shared enforcement',['Identity and office scope | Permissions | Input validation', 'Workflow prerequisites | Record version | Operation identifier'])
box((40,990,1760,1145),'Laravel application services',['Requests and clarification | Priority and assignment | Tracking', 'Feedback and ratings | Reporting | Audit and notifications'])
box((40,1250,580,1435),'MySQL',['Operational records', 'Audit and agent operations'])
box((650,1250,1190,1435),'Private storage',['Attachments', 'Generated reports'])
box((1260,1250,1760,1435),'Queue and scheduler',['Reports and notifications', 'Retries and job status'])
arrow([(345,195),(345,290)],(250,240,'HTTPS'))
arrow([(990,195),(990,290)],(1080,240,'Chat'))
arrow([(1475,195),(1475,245),(1280,245),(1280,530)])
arrow([(1240,365),(1330,365)])
arrow([(990,440),(990,530)])
arrow([(345,440),(345,755)])
arrow([(1240,660),(1240,755)])
arrow([(900,900),(900,990)])
arrow([(310,1145),(310,1250)])
arrow([(920,1145),(920,1250)])
arrow([(1510,1145),(1510,1250)])
im.save(QA/'architecture.png')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(25)
doc.styles['Subtitle'].font.size=Pt(12)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(5)
doc.styles['Heading 2'].font.size=Pt(11.5)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next=True

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('CLSU System Architecture  |  '); r.font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

def p(text, style=None): return doc.add_paragraph(text,style)
def h(text): doc.add_heading(text,2)
def page(title):
    doc.add_page_break(); doc.add_heading(title,1)
def bullets(items):
    for item in items: p(item,'List Bullet')
def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,text,w in zip(t.rows[0].cells,headers,widths): c.text=text; c.width=Inches(w)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,text,w in zip(t.add_row().cells,row,widths): c.text=text; c.width=Inches(w)
    for i,row in enumerate(t.rows):
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'243F52' if i==0 else ('F1F4F6' if i%2==0 else 'FFFFFF')); pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.size=Pt(9)
                    if i==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
    p('').paragraph_format.space_after=Pt(0)
    return t

doc.add_heading('CLSU Feedback Management System Architecture',0)
p('System architecture documentation','Subtitle')
p('Version 1.0  |  19 October 2026  |  Proposed implementation baseline')
h('Purpose and architectural decision')
p('This document defines the proposed architecture of the CLSU Feedback Management System and Service Tracker. It guides the capstone development team and MISO reviewers in implementing service requests, feedback collection, tracking, reporting, and controlled AI agent actions.')
p('The system will use one React application, one modular Laravel backend, and a centralized MySQL database. Both the web API and Model Context Protocol tools will invoke shared application services. This arrangement keeps business rules consistent and supports a manageable deployment for the development team.')
h('System scope')
p('The first release covers public general feedback, authenticated service requests, MISO review and assignment, personnel progress updates, clarification, completed-service evaluations, notifications, and standard reports. An MCP connection will let authorized administrators ask an AI agent what needs attention and instruct it to perform defined operational updates.')
table(['Layer','Selected approach'],[
('Frontend','React with TypeScript and Vite; responsive role-based interfaces'),
('Application','Laravel API, policies, validation, domain services and MCP server'),
('Persistence','MySQL and private attachment and report storage'),
('Background work','Laravel queue worker and scheduler'),
('AI integration','Compatible external MCP agent; server-side runtime for embedded chat')],[1.35,5.35])
h('Design basis')
p('The architecture follows the service-job and feedback requirements in the ITRM 3100 research draft and the requester, mobile, administrative, and tentative assistant screens in the FEEDBACK MANAGEMENT Figma file. The first-release MCP boundaries reflect the agreed capability scope. The architecture describes intended behavior and does not claim that the system has already been implemented or tested.')

page('1 Architecture overview')
p('The application uses a modular monolith: Laravel modules share a deployment and database while keeping their responsibilities separate. React communicates with the application API over HTTPS. External AI agents connect to a protected MCP endpoint; embedded chat uses a server-side agent runtime.')
doc.add_picture(str(QA/'architecture.png'),width=Inches(6.5))
p('Figure 1  Proposed logical system architecture').paragraph_format.space_after=Pt(10)
h('Common execution path')
p('API controllers and MCP tools call the same application services. Every protected operation resolves the caller identity, checks permissions and office scope, validates inputs, and applies workflow rules. Public feedback uses a separately restricted submission route. Operational changes and their audit entries are committed together where possible.')
h('Agent and frontend separation')
p('The model selects among exposed tools and explains returned results. Laravel determines whether an action is allowed and persists successful changes. The agent receives a structured result, while React reloads the affected record or responds to an application event. Model credentials remain on the server or in the external agent host.')

page('2 Application modules and user access')
table(['Module','Responsibility'],[
('Identity and access','Accounts, roles, office memberships and policy checks'),
('Requests and clarification','Unique job numbers, service details, revisions and clarification history'),
('Priority and assignment','Urgency and impact assessment, official priority, eligible staff and assignments'),
('Tracking','Permitted transitions, actions taken, timestamps and note visibility'),
('Feedback and ratings','General feedback and one requester evaluation per completed job'),
('Reporting','Consistent aggregates, period filters and private report exports'),
('Notifications and audit','Application notices and traceable changes from web and MCP actions')],[1.65,5.05])
h('Access model')
table(['Actor','Allowed access'],[
('Public visitor','Submit general feedback; no private request lookup by job number'),
('Authenticated requester','Create requests; view own records; clarify and rate eligible completed jobs'),
('Service personnel','View assigned work and perform permitted progress updates'),
('MISO administrator','Review and route requests within authorized scope; assign, report and use MCP'),
('Connected AI agent','Use approved tools within the connected administrator permissions and tool scopes')],[1.65,5.05])
h('Office and role boundaries')
p('The requesting office and handling office are separate fields. Office membership does not automatically grant access to all university records. Student, faculty, staff, visitor, and alumni classifications describe a user; operational roles determine permissions. The first release exposes operational MCP tools to authorized MISO administrators.')
h('Visibility of service records')
p('Requester-visible updates and internal operational notes are explicitly distinguished. Internal notes are excluded from requester responses. Private attachments require record-level authorization even when a user knows the file identifier or job number.')

page('3 Service workflows and business rules')
h('Request submission and review')
p('An authenticated requester submits the service type, subject, location, requesting affiliation, concern, urgency, and impact. Laravel validates the submission and generates a unique job number. MISO verifies the details, confirms the handling office and category, calculates priority using configured rules, and assigns eligible personnel.')
h('Clarification and progress')
p('MISO can request missing details and place the request in For Clarification. The requester resubmits under the same job number; revisions and messages remain in the history. Personnel record work and permitted status changes. On Hold requires a reason; resuming follows the configured transition rules.')
table(['Status','Meaning or prerequisite'],[
('Pending','Received and awaiting review or assignment'),
('For Clarification','A recorded clarification request needs a requester response'),
('Assigned','Eligible personnel are assigned to the request'),
('In Progress','Work has started and an operational update is recorded'),
('On Hold','Work is paused with a recorded reason'),
('Completed','Authorized staff record actions taken and the service outcome')],[1.45,5.25])
h('Completion and evaluation')
p('Authorized human staff mark work Completed after documenting the outcome. The first-release agent cannot complete a service. The requester may submit one optional evaluation linked to the job and personnel; not receiving a rating does not undo completion.')
h('General feedback')
p('General feedback can be submitted independently of a service job. Optional contact details allow follow-up. A supplied job number is an unverified reference until ownership or a secure invitation is verified; it grants neither record access nor permission to submit the verified job evaluation.')
h('Priority and measurement')
p('Requester-reported urgency and impact are retained separately from the official priority. Laravel applies the MISO-approved matrix; authorized exceptions require a reason. Reports distinguish first response time from completion time, define the reporting period and denominator, and show the rating count beside satisfaction averages. No response targets are assumed until MISO configures them.')

page('4 Data architecture and record integrity')
table(['Entity group','Main records and relationships'],[
('Access and organization','Users, offices, memberships, roles and permissions'),
('Service configuration','Service types, categories and versioned priority rules'),
('Service requests','Requester, requesting office, handling office, job number, details, official priority and status'),
('Operational history','Multiple assignments, status transitions, clarification messages, revisions and progress notes per request'),
('Feedback','Independent general feedback; zero or one verified requester evaluation per completed job'),
('Supporting records','Attachment metadata, notifications, report jobs, agent operations and append-only audit entries')],[1.65,5.05])
h('Database constraints')
p('Foreign keys preserve valid relationships. Unique constraints enforce job-number uniqueness, one verified evaluation per job, and operation-identifier uniqueness within the appropriate caller and tool scope. Job-number generation uses transactional sequencing or an equivalent concurrency-safe strategy rather than counting existing rows.')
h('Atomic updates and concurrent editing')
p('Each operational write checks an expected record version. If another user has changed the record, Laravel rejects the stale action and returns the current version for review. Related state changes and audit entries are saved in a database transaction. An assignment also checks personnel eligibility and current request state.')
h('Retry and notification behavior')
p('Every action tool accepts a unique operation identifier. Repeating the same operation returns its recorded result; reusing an identifier with different parameters is rejected. Notifications and exports are dispatched after commit. Workers also prevent duplicate delivery where feasible, since a database transaction cannot guarantee exactly-once delivery to an external email service.')
h('Storage and lifecycle')
p('MySQL stores attachment metadata; file content and generated reports remain in private storage. Downloads use authenticated authorization checks or short-lived authorized links. The implementation must define retention, backup frequency, restoration ownership, and deletion procedures with MISO before production use. The MCP tool set does not expose deletion.')

page('5 MCP tools and agent interaction')
p('MCP exposes specific business capabilities to compatible AI hosts. It supplies a tool interface; the connected host runs the agent and model interaction. The embedded assistant additionally requires a server-side runtime and configured provider. The initial tool catalog is shown below.')
table(['Tool','Purpose and limit'],[
('search_requests','Filter and paginate authorized requests; no arbitrary SQL'),
('get_request_details','Return permitted details and timeline with record version'),
('get_attention_queue','Return rule-based ranking, evidence and configured deadlines'),
('get_personnel_workload','Return recorded eligibility, availability and active assignments'),
('get_feedback_metrics','Return backend-calculated aggregates, periods and sample counts'),
('get_operation_result','Retrieve the caller-authorized outcome after a timeout or retry'),
('assign_request','Assign or reassign eligible personnel on explicit instruction'),
('set_request_priority','Apply rule-based priority or authorized documented exception'),
('add_internal_note','Append an explicitly requested operational note'),
('request_clarification','Record specified message and permitted status change'),
('transition_request_status','Apply configured operational transitions; exclude Completed'),
('request_report and get_report_job','Create a standard export and check its final outcome')],[2.5,4.2])
h('Recommendation followed by action')
p('For a question such as Which requests should we work on first, the agent reads the attention queue and recorded workload, then explains its recommendation with supporting job references. Reading and recommending do not authorize updates. An instruction such as Assign this request to Mark authorizes that defined action, subject to server checks.')
h('Tool contract')
p('Reads use bounded filters, pagination and minimal fields. Writes include the target identifier, validated parameters, expected version, operation identifier and reason. The server derives identity from authentication. Results contain the operation identifier, affected record, resulting version and success or error state. An accepted report job is not reported as a completed export.')
p('Permission denial, invalid workflow, stale version, invalid input and temporary service failure return distinct errors. The agent reports actual outcomes and does not claim success after a failed or merely queued operation.')

page('6 Security boundaries and agent limitations')
h('Authentication and authorization')
p('Protect the remote MCP endpoint with HTTPS and a supported authentication method. A controlled pilot can use revocable, scoped user tokens; external multi-user connections should use a compatible OAuth flow. Check roles, tool scopes, office scope and target-record permission on every call. A model-supplied user identifier cannot establish authority.')
h('Authorization of actions')
p('Explicit administrator instructions authorize the specified operational update. Where a separate confirmation is required, Laravel verifies an approval record bound to the user, exact target, parameters, record version and expiry. Approval is consumed once. Host approval prompts and tool annotations do not replace server authorization. Instructions embedded in service descriptions or feedback cannot authorize actions.')
h('First release exclusions')
bullets([
    'No arbitrary SQL, shell execution, source-code changes, or unrestricted database access.',
    'No account, role, permission, or system-configuration administration through MCP.',
    'No deletion or alteration of original requester statements, client ratings, or audit history.',
    'No agent-driven completion, fabricated service outcomes, or unattended bulk reassignment.',
    'No arbitrary outbound email, unrestricted exports, or continuous autonomous monitoring.'
])
h('Privacy and untrusted content')
p('Return only fields needed for the task and exclude passwords, tokens and unnecessary contact information. Validate attachments and public input, apply submission and tool rate limits, and separate private notes from public updates. Treat all submitted content and retrieved attachments as untrusted data. MISO must approve provider data handling and retention settings before personal data is processed by an external model.')
h('What AI recommendations can establish')
p('Recommendations use recorded urgency, impact, request age, configured deadlines and workload. Negative sentiment alone does not establish urgency. Recorded availability does not prove that a technician is physically available. The agent can explain trends and options, but cannot guarantee diagnosis accuracy or replace staff verification.')

page('7 Deployment and failure handling')
h('Initial deployment topology')
p('Deploy the React production build and Laravel API under one HTTPS domain where practical. Laravel connects to MySQL and private file storage; a supervised queue worker handles exports and notifications, and a scheduler handles defined maintenance tasks. Start with a database-backed queue if suitable for measured load; Redis is optional rather than a first-release requirement.')
p('A private campus deployment requires network access from the chosen agent host. An external MCP client cannot reach a campus-only endpoint without an approved network path. The AI provider also requires outbound access when used. Package versions, PHP extensions and client transport compatibility must be verified before deployment.')
h('Frontend synchronization')
p('After a successful mutation, React refetches the affected data. Polling can refresh other open views in the initial release; broadcast events may be added when needed. Persisting a database change does not automatically refresh a browser. A visible result should distinguish successful changes from queued, failed or conflicted operations.')
h('Failures and recovery')
table(['Condition','Required behavior'],[
('Model or MCP unavailable','Keep ordinary request, assignment, tracking and feedback functions available'),
('Tool timeout after submission','Retrieve the operation result using its identifier before retrying'),
('Concurrent record change','Reject stale write; refresh details and re-evaluate the action'),
('Report still running','Return a job identifier and status; provide download only after success'),
('Worker or email failure','Record failure, apply bounded retries and make operational failure visible')],[1.8,4.9])
h('Operational controls')
p('Set connection and request timeouts, worker supervision, retry limits and job monitoring. Store provider credentials and signing secrets outside the frontend and repository. Log failures without exposing sensitive request content or credentials. Test database and private-file restoration together.')
h('Capacity and availability')
p('The architecture does not assume unlimited throughput or continuous agent execution. Rate limits, model costs, context limits, host timeouts and university network conditions constrain use. Load testing will establish practical capacity; the research draft server specification is a planning baseline rather than a validated capacity claim.')

page('8 Verification and implementation decisions')
h('Acceptance criteria')
table(['Area','Required evidence before release'],[
('Request lifecycle','A requester submits, clarifies, tracks and optionally evaluates a completed job'),
('Permissions','Cross-user and cross-office access attempts fail through both API and MCP'),
('Priority and reports','Configured matrix, time calculations, denominators and rating counts produce expected results'),
('Agent actions','Authorized assignment and progress changes persist; excluded operations are unavailable'),
('Integrity','Stale versions fail; repeated operation identifiers do not duplicate writes'),
('Privacy','Internal notes and attachments remain restricted; injected instructions cannot grant access'),
('Resilience','Provider failure leaves core workflows usable; report jobs and failed notifications remain traceable')],[1.4,5.3])
h('Decisions to confirm during implementation')
p('MISO must confirm public feedback eligibility, guest access and moderation; the urgency and impact matrix and response targets; handling-office responsibilities; allowed status transitions; report period and business-hour definitions; provider data handling; retention and backup policies; and production hosting. These settings must be finalized before they become enforced business rules.')
p('The research manuscript must also align its scope with the inclusion of third-party AI integration and distinguish MCP from the language model. Exact Laravel, PHP, MCP and AI runtime versions will be pinned after compatibility checks.')
h('References')
p('Research draft: Development and Enhancement of Web Based Feedback Management System and Tracker for CLSU Services, ITRM 3100.md. Local reference supplied by the project team.')
p('Design reference: FEEDBACK MANAGEMENT, Figma. https://www.figma.com/design/PeQN6nVBQbPNXWoTU32Mhm/FEEDBACK-MANAGEMENT?node-id=0-1')
p('Laravel MCP documentation. https://laravel.com/framework/docs/13.x/mcp')
p('Laravel AI SDK documentation. https://laravel.com/framework/docs/13.x/ai-sdk')
p('Laravel authorization documentation. https://laravel.com/framework/docs/13.x/authorization')
p('Laravel queues documentation. https://laravel.com/framework/docs/13.x/queues')
p('Documentation references reviewed 19 October 2026. Framework features are distinguished from the application rules and controls specified in this document.')

doc.core_properties.title='CLSU Feedback Management System Architecture'
doc.core_properties.subject='System architecture and MCP capability specification'
doc.core_properties.author='CLSU Feedback Management Project'
doc.core_properties.keywords='React, Laravel, MySQL, MCP, service requests, feedback'
# Remove inherited Word title rules as well as any direct paragraph borders.
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for shape in doc.inline_shapes:
    shape._inline.docPr.set('descr','Architecture showing React and external AI clients entering Laravel through API or MCP tools, shared controls and services, and MySQL, private storage and background jobs.')
doc.save(OUT)
print(OUT)
